from __future__ import annotations

import hashlib
from datetime import date, datetime
from typing import TYPE_CHECKING

from werkzeug.utils import secure_filename

from app.eqms.audit import record_event
from app.eqms.utils import utcnow

if TYPE_CHECKING:
    from sqlalchemy.orm import Session
    from app.eqms.models import User
    from app.eqms.modules.equipment.models import Equipment, EquipmentSupplier, ManagedDocument
    from app.eqms.modules.suppliers.models import Supplier


VALID_STATUSES = ("Active", "Inactive", "Retired", "Calibration Overdue", "PM Overdue")

# Days-out threshold under which a due date is flagged "Due soon".
DUE_SOON_DAYS = 30


def due_status(due: date | None, interval_text: str | None = None, today: date | None = None) -> dict:
    """
    Compute an at-a-glance calibration/PM status from a due date.

    Returns a dict: {"state", "label", "days"}. States:
      - "none":     no due date and no meaningful interval (e.g. "N/A") -> not tracked
      - "unscheduled": interval is set (e.g. "Annual") but no due date recorded
      - "overdue":  due date is in the past
      - "due_soon": due within DUE_SOON_DAYS
      - "ok":       due further out
    """
    if today is None:
        today = date.today()
    it = (interval_text or "").strip()
    it_is_na = it == "" or it.upper() in ("N/A", "NA", "NONE")
    if due is None:
        if it_is_na:
            return {"state": "none", "label": "Not tracked", "days": None}
        return {"state": "unscheduled", "label": "No date on file", "days": None}
    days = (due - today).days
    if days < 0:
        return {"state": "overdue", "label": f"Overdue {abs(days)}d", "days": days}
    if days <= DUE_SOON_DAYS:
        return {"state": "due_soon", "label": f"Due in {days}d", "days": days}
    return {"state": "ok", "label": f"Due {due.isoformat()}", "days": days}


def parse_date(s: str | None) -> date | None:
    """Parse YYYY-MM-DD date string."""
    if not s:
        return None
    s = s.strip()
    if not s:
        return None
    return date.fromisoformat(s)


def coerce_cell_date(value) -> date | None:
    """
    Coerce a spreadsheet cell into a date. Accepts datetime/date objects and a
    range of human date strings used in the SILQ master lists ("11/1/22",
    "12 DEC 2025", "2026-11-01"). Returns None for blanks / "N/A".
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.upper() in ("N/A", "NA", "NONE", "-"):
        return None
    # ISO first
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        pass
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%d %b %Y", "%d%b%Y", "%d %B %Y", "%b %d %Y", "%B %d %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    return text or None


# Header labels in the SILQ Equipment Master List, mapped to payload keys.
_EQUIP_MASTER_COLUMNS = {
    "equip id": "equip_code",
    "equip status": "status",
    "equipment description": "description",
    "mfg": "mfg",
    "model no.": "model_no",
    "model no": "model_no",
    "serial no.": "serial_no",
    "serial no": "serial_no",
    "date put in-service": "date_in_service",
    "location": "location",
    "cal interval": "cal_interval_text",
    "date of last cal": "last_cal_date",
    "cal due date": "cal_due_date",
    "pm interval": "pm_interval_text",
    "date of last pm": "last_pm_date",
    "pm due date": "pm_due_date",
    "comments": "comments",
}

_DATE_KEYS = {"date_in_service", "last_cal_date", "cal_due_date", "last_pm_date", "pm_due_date"}


def import_equipment_master(s: "Session", file_bytes: bytes, user: "User") -> dict:
    """
    Upsert equipment from an uploaded SILQ Equipment Master List (.xlsx).

    Matches the header row wherever it appears, keys rows by Equip ID, and
    creates or updates each row. Free-text intervals are stored verbatim; dates
    are coerced leniently. Returns a summary dict.
    """
    import io
    import openpyxl

    from app.eqms.modules.equipment.models import Equipment

    result = {"created": 0, "updated": 0, "skipped": 0, "errors": []}
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Locate the header row (contains "Equip ID" / "Equip\nID").
    header_idx = None
    col_map: dict[int, str] = {}
    for i, row in enumerate(rows):
        norm = [(" ".join(str(c).split())).strip().lower() if c is not None else "" for c in row]
        if any(v in ("equip id",) for v in norm):
            for j, label in enumerate(norm):
                if label in _EQUIP_MASTER_COLUMNS:
                    col_map[j] = _EQUIP_MASTER_COLUMNS[label]
            header_idx = i
            break
    if header_idx is None or "equip_code" not in col_map.values():
        result["errors"].append("Could not find the Equipment Master List header row (expected an 'Equip ID' column).")
        return result

    for row in rows[header_idx + 1:]:
        payload: dict = {}
        for j, key in col_map.items():
            value = row[j] if j < len(row) else None
            if key in _DATE_KEYS:
                d = coerce_cell_date(value)
                payload[key] = d.isoformat() if d else None
            else:
                payload[key] = _cell_text(value)
        code = (payload.get("equip_code") or "").strip()
        if not code:
            continue
        status = (payload.get("status") or "").strip()
        if status not in VALID_STATUSES:
            payload["status"] = "Active"
        existing = s.query(Equipment).filter(Equipment.equip_code == code).one_or_none()
        try:
            if existing:
                update_equipment(s, existing, payload, user, reason="Equipment Master List import")
                result["updated"] += 1
            else:
                create_equipment(s, payload, user)
                result["created"] += 1
        except Exception as e:  # noqa: BLE001 - collect per-row errors, keep importing
            result["errors"].append(f"{code}: {e}")
            result["skipped"] += 1
    return result


def parse_int(s: str | None) -> int | None:
    """Parse integer string."""
    if not s:
        return None
    s = s.strip()
    if not s:
        return None
    return int(s)


def validate_equipment_payload(payload: dict) -> list[str]:
    """Validate equipment creation/update payload. Returns list of errors."""
    errors = []
    equip_code = (payload.get("equip_code") or "").strip()
    if not equip_code:
        errors.append("Equipment code is required.")
    status = (payload.get("status") or "").strip()
    if status and status not in VALID_STATUSES:
        errors.append(f"Invalid status. Must be one of: {', '.join(VALID_STATUSES)}")
    return errors


def create_equipment(s: "Session", payload: dict, user: "User") -> "Equipment":
    """Create new equipment."""
    from app.eqms.modules.equipment.models import Equipment

    now = utcnow()
    equipment = Equipment(
        equip_code=(payload.get("equip_code") or "").strip(),
        status=(payload.get("status") or "Active").strip(),
        description=(payload.get("description") or "").strip() or None,
        mfg=(payload.get("mfg") or "").strip() or None,
        model_no=(payload.get("model_no") or "").strip() or None,
        serial_no=(payload.get("serial_no") or "").strip() or None,
        date_in_service=parse_date(payload.get("date_in_service")),
        location=(payload.get("location") or "").strip() or None,
        cal_interval=parse_int(payload.get("cal_interval")),
        cal_interval_text=(payload.get("cal_interval_text") or "").strip() or None,
        last_cal_date=parse_date(payload.get("last_cal_date")),
        cal_due_date=parse_date(payload.get("cal_due_date")),
        pm_interval=parse_int(payload.get("pm_interval")),
        pm_interval_text=(payload.get("pm_interval_text") or "").strip() or None,
        last_pm_date=parse_date(payload.get("last_pm_date")),
        pm_due_date=parse_date(payload.get("pm_due_date")),
        comments=(payload.get("comments") or "").strip() or None,
        custom_fields=payload.get("custom_fields") if isinstance(payload.get("custom_fields"), dict) else None,
        created_at=now,
        updated_at=now,
        created_by_user_id=user.id,
        updated_by_user_id=user.id,
    )
    s.add(equipment)
    s.flush()

    record_event(
        s,
        actor=user,
        action="equipment.create",
        entity_type="Equipment",
        entity_id=str(equipment.id),
        metadata={"equip_code": equipment.equip_code, "status": equipment.status, "description": equipment.description},
    )
    return equipment


def update_equipment(s: "Session", equipment: "Equipment", payload: dict, user: "User", reason: str | None = None) -> "Equipment":
    """Update existing equipment."""
    changes = {}

    # equip_code is read-only after creation, but track if provided
    new_status = (payload.get("status") or "").strip()
    if new_status and new_status != equipment.status:
        changes["status"] = {"old": equipment.status, "new": new_status}
        equipment.status = new_status

    new_description = (payload.get("description") or "").strip() or None
    if new_description != equipment.description:
        changes["description"] = {"old": equipment.description, "new": new_description}
        equipment.description = new_description

    new_mfg = (payload.get("mfg") or "").strip() or None
    if new_mfg != equipment.mfg:
        changes["mfg"] = {"old": equipment.mfg, "new": new_mfg}
        equipment.mfg = new_mfg

    new_model_no = (payload.get("model_no") or "").strip() or None
    if new_model_no != equipment.model_no:
        changes["model_no"] = {"old": equipment.model_no, "new": new_model_no}
        equipment.model_no = new_model_no

    new_serial_no = (payload.get("serial_no") or "").strip() or None
    if new_serial_no != equipment.serial_no:
        changes["serial_no"] = {"old": equipment.serial_no, "new": new_serial_no}
        equipment.serial_no = new_serial_no

    new_date_in_service = parse_date(payload.get("date_in_service"))
    if new_date_in_service != equipment.date_in_service:
        changes["date_in_service"] = {"old": str(equipment.date_in_service), "new": str(new_date_in_service)}
        equipment.date_in_service = new_date_in_service

    new_location = (payload.get("location") or "").strip() or None
    if new_location != equipment.location:
        changes["location"] = {"old": equipment.location, "new": new_location}
        equipment.location = new_location

    new_cal_interval = parse_int(payload.get("cal_interval"))
    if new_cal_interval != equipment.cal_interval:
        changes["cal_interval"] = {"old": equipment.cal_interval, "new": new_cal_interval}
        equipment.cal_interval = new_cal_interval

    new_cal_interval_text = (payload.get("cal_interval_text") or "").strip() or None
    if new_cal_interval_text != equipment.cal_interval_text:
        changes["cal_interval_text"] = {"old": equipment.cal_interval_text, "new": new_cal_interval_text}
        equipment.cal_interval_text = new_cal_interval_text

    new_last_cal_date = parse_date(payload.get("last_cal_date"))
    if new_last_cal_date != equipment.last_cal_date:
        changes["last_cal_date"] = {"old": str(equipment.last_cal_date), "new": str(new_last_cal_date)}
        equipment.last_cal_date = new_last_cal_date

    new_cal_due_date = parse_date(payload.get("cal_due_date"))
    if new_cal_due_date != equipment.cal_due_date:
        changes["cal_due_date"] = {"old": str(equipment.cal_due_date), "new": str(new_cal_due_date)}
        equipment.cal_due_date = new_cal_due_date

    new_pm_interval = parse_int(payload.get("pm_interval"))
    if new_pm_interval != equipment.pm_interval:
        changes["pm_interval"] = {"old": equipment.pm_interval, "new": new_pm_interval}
        equipment.pm_interval = new_pm_interval

    new_pm_interval_text = (payload.get("pm_interval_text") or "").strip() or None
    if new_pm_interval_text != equipment.pm_interval_text:
        changes["pm_interval_text"] = {"old": equipment.pm_interval_text, "new": new_pm_interval_text}
        equipment.pm_interval_text = new_pm_interval_text

    new_last_pm_date = parse_date(payload.get("last_pm_date"))
    if new_last_pm_date != equipment.last_pm_date:
        changes["last_pm_date"] = {"old": str(equipment.last_pm_date), "new": str(new_last_pm_date)}
        equipment.last_pm_date = new_last_pm_date

    new_pm_due_date = parse_date(payload.get("pm_due_date"))
    if new_pm_due_date != equipment.pm_due_date:
        changes["pm_due_date"] = {"old": str(equipment.pm_due_date), "new": str(new_pm_due_date)}
        equipment.pm_due_date = new_pm_due_date

    new_comments = (payload.get("comments") or "").strip() or None
    if new_comments != equipment.comments:
        changes["comments"] = {"old": equipment.comments, "new": new_comments}
        equipment.comments = new_comments

    new_custom_fields = payload.get("custom_fields") if isinstance(payload.get("custom_fields"), dict) else None
    if new_custom_fields is not None and new_custom_fields != (equipment.custom_fields or {}):
        changes["custom_fields"] = {"old": equipment.custom_fields or {}, "new": new_custom_fields}
        equipment.custom_fields = new_custom_fields

    equipment.updated_at = utcnow()
    equipment.updated_by_user_id = user.id

    record_event(
        s,
        actor=user,
        action="equipment.edit",
        entity_type="Equipment",
        entity_id=str(equipment.id),
        reason=reason,
        metadata={"equip_code": equipment.equip_code, "changes": changes},
    )
    return equipment


def build_equipment_storage_key(equip_code: str, filename: str, upload_date: date | None = None) -> str:
    """Build deterministic storage key for equipment document."""
    if upload_date is None:
        upload_date = date.today()
    safe_code = equip_code.replace("/", "_").replace("\\", "_")  # Prevent path traversal
    safe_filename = secure_filename(filename) or "document.bin"
    return f"equipment/{safe_code}/{upload_date.isoformat()}/{safe_filename}"


def file_digest_and_bytes(file_bytes: bytes) -> tuple[str, int]:
    """Compute SHA256 digest and size."""
    h = hashlib.sha256()
    h.update(file_bytes)
    return (h.hexdigest(), len(file_bytes))


def upload_equipment_document(
    s: "Session",
    equipment: "Equipment",
    file_bytes: bytes,
    filename: str,
    content_type: str,
    user: "User",
    description: str | None = None,
    document_type: str | None = None,
    extracted_text: str | None = None,
    category: str | None = None,
    is_primary: bool = False,
) -> "ManagedDocument":
    """Upload a document to equipment."""
    from flask import current_app
    from app.eqms.storage import storage_from_config
    from app.eqms.modules.equipment.models import ManagedDocument
    from app.eqms.utils import validate_managed_document

    sha256, size_bytes = file_digest_and_bytes(file_bytes)
    storage_key = build_equipment_storage_key(equipment.equip_code, filename)

    storage = storage_from_config(current_app.config)
    storage.put_bytes(storage_key, file_bytes, content_type=content_type)

    doc = ManagedDocument(
        entity_type="equipment",
        entity_id=equipment.id,
        equipment_id=equipment.id,
        storage_key=storage_key,
        original_filename=secure_filename(filename) or "document.bin",
        content_type=content_type,
        sha256=sha256,
        size_bytes=size_bytes,
        description=description,
        document_type=document_type,
        extracted_text=extracted_text,
        category=category or "general",
        is_primary=bool(is_primary),
        uploaded_by_user_id=user.id,
    )
    validate_managed_document(doc)
    s.add(doc)
    s.flush()

    record_event(
        s,
        actor=user,
        action="equipment.document_upload",
        entity_type="ManagedDocument",
        entity_id=str(doc.id),
        metadata={
            "equipment_id": equipment.id,
            "equip_code": equipment.equip_code,
            "filename": doc.original_filename,
            "document_type": document_type,
            "category": doc.category,
            "is_primary": doc.is_primary,
        },
    )
    return doc


def delete_equipment_document(s: "Session", document: "ManagedDocument", user: "User", reason: str) -> None:
    """Soft-delete an equipment document."""
    document.is_deleted = True
    document.deleted_at = utcnow()
    document.deleted_by_user_id = user.id

    record_event(
        s,
        actor=user,
        action="equipment.document_delete",
        entity_type="ManagedDocument",
        entity_id=str(document.id),
        reason=reason,
        metadata={
            "equipment_id": document.equipment_id,
            "filename": document.original_filename,
        },
    )


def add_supplier_to_equipment(
    s: "Session",
    equipment: "Equipment",
    supplier: "Supplier",
    relationship_type: str | None,
    notes: str | None,
    user: "User",
) -> "EquipmentSupplier":
    """Add supplier association to equipment."""
    from app.eqms.modules.equipment.models import EquipmentSupplier

    assoc = EquipmentSupplier(
        equipment_id=equipment.id,
        supplier_id=supplier.id,
        relationship_type=relationship_type,
        notes=notes,
        created_by_user_id=user.id,
    )
    s.add(assoc)
    s.flush()

    record_event(
        s,
        actor=user,
        action="equipment.supplier_added",
        entity_type="EquipmentSupplier",
        entity_id=str(assoc.id),
        metadata={
            "equipment_id": equipment.id,
            "supplier_id": supplier.id,
            "relationship_type": relationship_type,
        },
    )
    return assoc


def remove_supplier_from_equipment(s: "Session", association: "EquipmentSupplier", user: "User", reason: str) -> None:
    """Remove supplier association from equipment."""
    equipment_id = association.equipment_id
    supplier_id = association.supplier_id
    assoc_id = association.id

    record_event(
        s,
        actor=user,
        action="equipment.supplier_removed",
        entity_type="EquipmentSupplier",
        entity_id=str(assoc_id),
        reason=reason,
        metadata={
            "equipment_id": equipment_id,
            "supplier_id": supplier_id,
        },
    )
    s.delete(association)
