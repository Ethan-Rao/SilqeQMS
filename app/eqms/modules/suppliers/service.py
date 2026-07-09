from __future__ import annotations

import hashlib
from datetime import date, datetime
from typing import TYPE_CHECKING

from werkzeug.utils import secure_filename

from app.eqms.audit import record_event
from app.eqms.utils import utcnow

if TYPE_CHECKING:
    from sqlalchemy.orm import Session
    from app.eqms.models import User
    from app.eqms.modules.suppliers.models import Supplier
    from app.eqms.utils import validate_managed_document
    from app.eqms.modules.equipment.models import ManagedDocument


VALID_STATUSES = ("Approved", "Conditional", "Pending", "Rejected")

# Days-out threshold under which a re-evaluation / certification date is "due soon".
DUE_SOON_DAYS = 60


def date_status(due: date | None, today: date | None = None) -> dict:
    """
    At-a-glance status for a re-evaluation or certification expiration date.
    States: "none" (no date), "overdue", "due_soon", "ok".
    """
    if today is None:
        today = date.today()
    if due is None:
        return {"state": "none", "label": "—", "days": None}
    days = (due - today).days
    if days < 0:
        return {"state": "overdue", "label": f"Overdue {abs(days)}d", "days": days}
    if days <= DUE_SOON_DAYS:
        return {"state": "due_soon", "label": f"{due.isoformat()} ({days}d)", "days": days}
    return {"state": "ok", "label": due.isoformat(), "days": days}


def parse_date(s: str | None) -> date | None:
    """Parse YYYY-MM-DD date string."""
    if not s:
        return None
    s = s.strip()
    if not s:
        return None
    return date.fromisoformat(s)


def validate_supplier_payload(payload: dict) -> list[str]:
    """Validate supplier creation/update payload. Returns list of errors."""
    errors = []
    name = (payload.get("name") or "").strip()
    if not name:
        errors.append("Name is required.")
    status = (payload.get("status") or "").strip()
    if status and status not in VALID_STATUSES:
        errors.append(f"Invalid status. Must be one of: {', '.join(VALID_STATUSES)}")
    return errors


def create_supplier(s: "Session", payload: dict, user: "User") -> "Supplier":
    """Create a new supplier."""
    from app.eqms.modules.suppliers.models import Supplier

    now = utcnow()
    supplier = Supplier(
        name=(payload.get("name") or "").strip(),
        status=(payload.get("status") or "Pending").strip(),
        category=(payload.get("category") or "").strip() or None,
        product_service_provided=(payload.get("product_service_provided") or "").strip() or None,
        address=(payload.get("address") or "").strip() or None,
        contact_name=(payload.get("contact_name") or "").strip() or None,
        contact_email=(payload.get("contact_email") or "").strip() or None,
        contact_phone=(payload.get("contact_phone") or "").strip() or None,
        initial_listing_date=parse_date(payload.get("initial_listing_date")),
        certification_expiration=parse_date(payload.get("certification_expiration")),
        certification_type=(payload.get("certification_type") or "").strip() or None,
        next_reevaluation_date=parse_date(payload.get("next_reevaluation_date")),
        notes=(payload.get("notes") or "").strip() or None,
        custom_fields=payload.get("custom_fields") if isinstance(payload.get("custom_fields"), dict) else None,
        created_at=now,
        updated_at=now,
        created_by_user_id=user.id,
        updated_by_user_id=user.id,
    )
    s.add(supplier)
    s.flush()

    record_event(
        s,
        actor=user,
        action="supplier.create",
        entity_type="Supplier",
        entity_id=str(supplier.id),
        metadata={"name": supplier.name, "status": supplier.status},
    )
    return supplier


def update_supplier(s: "Session", supplier: "Supplier", payload: dict, user: "User", reason: str | None = None) -> "Supplier":
    """Update an existing supplier."""
    changes = {}
    old_status = supplier.status

    new_name = (payload.get("name") or "").strip()
    if new_name and new_name != supplier.name:
        changes["name"] = {"old": supplier.name, "new": new_name}
        supplier.name = new_name

    new_status = (payload.get("status") or "").strip()
    if new_status and new_status != supplier.status:
        changes["status"] = {"old": supplier.status, "new": new_status}
        supplier.status = new_status

    new_category = (payload.get("category") or "").strip() or None
    if new_category != supplier.category:
        changes["category"] = {"old": supplier.category, "new": new_category}
        supplier.category = new_category

    new_psp = (payload.get("product_service_provided") or "").strip() or None
    if new_psp != supplier.product_service_provided:
        changes["product_service_provided"] = {"old": supplier.product_service_provided, "new": new_psp}
        supplier.product_service_provided = new_psp

    new_address = (payload.get("address") or "").strip() or None
    if new_address != supplier.address:
        changes["address"] = {"old": supplier.address, "new": new_address}
        supplier.address = new_address

    new_contact_name = (payload.get("contact_name") or "").strip() or None
    if new_contact_name != supplier.contact_name:
        changes["contact_name"] = {"old": supplier.contact_name, "new": new_contact_name}
        supplier.contact_name = new_contact_name

    new_contact_email = (payload.get("contact_email") or "").strip() or None
    if new_contact_email != supplier.contact_email:
        changes["contact_email"] = {"old": supplier.contact_email, "new": new_contact_email}
        supplier.contact_email = new_contact_email

    new_contact_phone = (payload.get("contact_phone") or "").strip() or None
    if new_contact_phone != supplier.contact_phone:
        changes["contact_phone"] = {"old": supplier.contact_phone, "new": new_contact_phone}
        supplier.contact_phone = new_contact_phone

    new_ild = parse_date(payload.get("initial_listing_date"))
    if new_ild != supplier.initial_listing_date:
        changes["initial_listing_date"] = {"old": str(supplier.initial_listing_date), "new": str(new_ild)}
        supplier.initial_listing_date = new_ild

    new_ce = parse_date(payload.get("certification_expiration"))
    if new_ce != supplier.certification_expiration:
        changes["certification_expiration"] = {"old": str(supplier.certification_expiration), "new": str(new_ce)}
        supplier.certification_expiration = new_ce

    new_ct = (payload.get("certification_type") or "").strip() or None
    if new_ct != supplier.certification_type:
        changes["certification_type"] = {"old": supplier.certification_type, "new": new_ct}
        supplier.certification_type = new_ct

    new_nrd = parse_date(payload.get("next_reevaluation_date"))
    if new_nrd != supplier.next_reevaluation_date:
        changes["next_reevaluation_date"] = {"old": str(supplier.next_reevaluation_date), "new": str(new_nrd)}
        supplier.next_reevaluation_date = new_nrd

    new_notes = (payload.get("notes") or "").strip() or None
    if new_notes != supplier.notes:
        changes["notes"] = {"old": supplier.notes, "new": new_notes}
        supplier.notes = new_notes

    new_custom_fields = payload.get("custom_fields") if isinstance(payload.get("custom_fields"), dict) else None
    if new_custom_fields is not None and new_custom_fields != (supplier.custom_fields or {}):
        changes["custom_fields"] = {"old": supplier.custom_fields or {}, "new": new_custom_fields}
        supplier.custom_fields = new_custom_fields

    supplier.updated_at = utcnow()
    supplier.updated_by_user_id = user.id

    record_event(
        s,
        actor=user,
        action="supplier.edit",
        entity_type="Supplier",
        entity_id=str(supplier.id),
        reason=reason,
        metadata={"name": supplier.name, "changes": changes},
    )
    return supplier


# Header labels accepted by the Approved Supplier List importer, mapped to payload keys.
_SUPPLIER_IMPORT_COLUMNS = {
    "supplier/contractor name": "name",
    "supplier": "name",
    "name": "name",
    "address": "address",
    "product/ service category": "category",
    "product/service category": "category",
    "category": "category",
    "product / service provided": "product_service_provided",
    "product/service provided": "product_service_provided",
    "product_service_provided": "product_service_provided",
    "scope": "product_service_provided",
    "initial listing date": "initial_listing_date",
    "initial_listing_date": "initial_listing_date",
    "status": "status",
    "notes / comments": "notes",
    "notes/comments": "notes",
    "notes": "notes",
    "certification type": "certification_type",
    "certification_type": "certification_type",
    "certification expiration": "certification_expiration",
    "certification_expiration": "certification_expiration",
    "next re-evaluation date": "next_reevaluation_date",
    "next reevaluation date": "next_reevaluation_date",
    "next_reevaluation_date": "next_reevaluation_date",
}

_SUPPLIER_DATE_KEYS = {"initial_listing_date", "certification_expiration", "next_reevaluation_date"}


def _coerce_import_date(value) -> str | None:
    """Coerce a cell/string into an ISO date string (or None)."""
    from app.eqms.modules.equipment.service import coerce_cell_date

    d = coerce_cell_date(value)
    return d.isoformat() if d else None


def import_supplier_list(s: "Session", file_bytes: bytes, filename: str, user: "User") -> dict:
    """
    Upsert suppliers from an uploaded Approved Supplier List (.xlsx or .csv).

    Requires a header row; matches columns leniently (see _SUPPLIER_IMPORT_COLUMNS)
    and keys rows by supplier name (case-insensitive). Returns a summary dict.
    """
    from app.eqms.modules.suppliers.models import Supplier

    result = {"created": 0, "updated": 0, "skipped": 0, "errors": []}

    # Read rows as a list of tuples regardless of format.
    rows: list[tuple] = []
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        import csv
        import io

        text = file_bytes.decode("utf-8-sig", errors="replace")
        rows = [tuple(r) for r in csv.reader(io.StringIO(text))]
    else:
        import io
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

    header_idx = None
    col_map: dict[int, str] = {}
    for i, row in enumerate(rows):
        norm = [(" ".join(str(c).split())).strip().lower() if c is not None else "" for c in row]
        if any(v in ("name", "supplier", "supplier/contractor name") for v in norm):
            for j, label in enumerate(norm):
                if label in _SUPPLIER_IMPORT_COLUMNS:
                    col_map[j] = _SUPPLIER_IMPORT_COLUMNS[label]
            header_idx = i
            break
    if header_idx is None or "name" not in col_map.values():
        result["errors"].append("Could not find a supplier header row (expected a 'Name' or 'Supplier' column).")
        return result

    for row in rows[header_idx + 1:]:
        payload: dict = {}
        for j, key in col_map.items():
            value = row[j] if j < len(row) else None
            if key in _SUPPLIER_DATE_KEYS:
                payload[key] = _coerce_import_date(value)
            else:
                text = None if value is None else str(value).strip()
                payload[key] = text or None
        name = (payload.get("name") or "").strip()
        if not name:
            continue
        status = (payload.get("status") or "").strip()
        # ASL "Notes/Comments" sometimes carries approval status words like "Approved"/"Conditional".
        if status not in VALID_STATUSES:
            note = (payload.get("notes") or "")
            for candidate in VALID_STATUSES:
                if candidate.lower() in note.lower():
                    status = candidate
                    break
            payload["status"] = status if status in VALID_STATUSES else "Approved"
        existing = (
            s.query(Supplier).filter(Supplier.name.ilike(name)).order_by(Supplier.id.asc()).first()
        )
        try:
            if existing:
                update_supplier(s, existing, payload, user, reason="Approved Supplier List import")
                result["updated"] += 1
            else:
                create_supplier(s, payload, user)
                result["created"] += 1
        except Exception as e:  # noqa: BLE001
            result["errors"].append(f"{name}: {e}")
            result["skipped"] += 1
    return result


def build_supplier_storage_key(supplier_id: int, filename: str, upload_date: date | None = None) -> str:
    """Build deterministic storage key for supplier document."""
    if upload_date is None:
        upload_date = date.today()
    safe_filename = secure_filename(filename) or "document.bin"
    return f"suppliers/{supplier_id}/{upload_date.isoformat()}/{safe_filename}"


def file_digest_and_bytes(file_bytes: bytes) -> tuple[str, int]:
    """Compute SHA256 digest and size."""
    h = hashlib.sha256()
    h.update(file_bytes)
    return (h.hexdigest(), len(file_bytes))


def upload_supplier_document(
    s: "Session",
    supplier: "Supplier",
    file_bytes: bytes,
    filename: str,
    content_type: str,
    user: "User",
    description: str | None = None,
    document_type: str | None = None,
    extracted_text: str | None = None,
) -> "ManagedDocument":
    """Upload a document to a supplier."""
    from flask import current_app
    from app.eqms.storage import storage_from_config
    from app.eqms.modules.equipment.models import ManagedDocument
    from app.eqms.utils import validate_managed_document  # runtime import (C-001)

    sha256, size_bytes = file_digest_and_bytes(file_bytes)
    storage_key = build_supplier_storage_key(supplier.id, filename)

    storage = storage_from_config(current_app.config)
    storage.put_bytes(storage_key, file_bytes, content_type=content_type)

    doc = ManagedDocument(
        entity_type="supplier",
        entity_id=supplier.id,
        supplier_id=supplier.id,
        storage_key=storage_key,
        original_filename=secure_filename(filename) or "document.bin",
        content_type=content_type,
        sha256=sha256,
        size_bytes=size_bytes,
        description=description,
        document_type=document_type,
        extracted_text=extracted_text,
        uploaded_by_user_id=user.id,
    )
    validate_managed_document(doc)
    s.add(doc)
    s.flush()

    record_event(
        s,
        actor=user,
        action="supplier.document_upload",
        entity_type="ManagedDocument",
        entity_id=str(doc.id),
        metadata={
            "supplier_id": supplier.id,
            "name": supplier.name,
            "filename": doc.original_filename,
            "document_type": document_type,
        },
    )
    return doc


def delete_supplier_document(s: "Session", document: "ManagedDocument", user: "User", reason: str) -> None:
    """Soft-delete a supplier document."""
    document.is_deleted = True
    document.deleted_at = utcnow()
    document.deleted_by_user_id = user.id

    record_event(
        s,
        actor=user,
        action="supplier.document_delete",
        entity_type="ManagedDocument",
        entity_id=str(document.id),
        reason=reason,
        metadata={
            "supplier_id": document.supplier_id,
            "filename": document.original_filename,
        },
    )
